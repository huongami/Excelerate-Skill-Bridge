"""Generate ten synthetic tech JDs using the verified Skill Bridge JD template."""

import importlib.util
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
BASE_GENERATOR = HERE.parent / "generate_demo_jds.py"
spec = importlib.util.spec_from_file_location("skill_bridge_jd_template", BASE_GENERATOR)
template = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(template)
template.OUT = HERE

TECH_JDS = [
    ("11", "Data Analyst", "Atlas Community Health", "Melbourne, VIC", [
        "SQL and data visualisation using Power BI are essential.",
        "You must demonstrate spreadsheet analysis and statistical analysis.",
        "Stakeholder communication and data-informed decisions are required.",
        "Python experience is preferred."],
     ["Query and validate operational datasets for recurring analysis.", "Build clear Power BI dashboards and explain trends to service leaders.", "Document metric definitions and investigate data-quality issues."]),
    ("12", "Data Engineer", "SignalWorks Digital", "Melbourne, VIC", [
        "Python, SQL and ETL pipelines are essential.",
        "You must have data warehousing and data modelling experience.",
        "Cloud platforms such as AWS or Azure are required.",
        "Data governance and CI/CD experience are preferred."],
     ["Build reliable batch and streaming data pipelines.", "Model curated warehouse tables for analytics consumers.", "Monitor pipeline quality, failures and recovery actions."]),
    ("13", "Analytics Engineer", "Northstar Marketplace", "Richmond, VIC", [
        "SQL, data modelling and data warehousing are essential.",
        "You must build tested ETL pipelines for analytics datasets.",
        "Data governance and stakeholder communication are required.",
        "Python and cloud platforms are preferred."],
     ["Transform raw warehouse data into trusted reporting models.", "Define tests, documentation and lineage for key datasets.", "Partner with analysts to improve reusable business metrics."]),
    ("14", "Backend Software Engineer", "HarbourTech Platforms", "Melbourne, VIC", [
        "Backend development and API development are essential.",
        "You must demonstrate database design and software testing.",
        "Cloud platforms and CI/CD are required.",
        "Container orchestration is preferred."],
     ["Design and maintain secure REST APIs and backend services.", "Write automated unit and integration tests.", "Review production issues and improve service reliability."]),
    ("15", "Frontend Software Engineer", "BrightLoop Products", "Carlton, VIC", [
        "Frontend development with React and TypeScript is essential.",
        "You must demonstrate software testing and API integration.",
        "Stakeholder communication and agile delivery are required.",
        "Next.js and CI/CD experience are preferred."],
     ["Build accessible, responsive product interfaces.", "Integrate frontend features with documented APIs.", "Contribute tests, code reviews and iterative product delivery."]),
    ("16", "DevOps Engineer", "CivicCloud Services", "Docklands, VIC", [
        "CI/CD, cloud platforms and infrastructure as code are essential.",
        "You must demonstrate container orchestration with Kubernetes and Docker.",
        "Risk and issue management and operational reporting are required.",
        "Python automation is preferred."],
     ["Build reliable deployment pipelines and cloud environments.", "Automate infrastructure changes using Terraform.", "Monitor incidents, document recovery actions and improve resilience."]),
    ("17", "Cybersecurity Analyst", "SecurePath Australia", "Melbourne, VIC", [
        "Cybersecurity monitoring and incident detection are essential.",
        "You must demonstrate identity and access management and risk management.",
        "Operational reporting and stakeholder communication are required.",
        "Python and cloud platforms are preferred."],
     ["Triage security alerts and document investigation evidence.", "Support access reviews and incident-response activities.", "Prepare clear security reports for technical and business stakeholders."]),
    ("18", "Machine Learning Engineer", "ModelWorks AI", "Southbank, VIC", [
        "Python and machine learning are essential.",
        "You must demonstrate MLOps, model deployment and software testing.",
        "Data engineering pipelines and cloud platforms are required.",
        "Container orchestration is preferred."],
     ["Develop and evaluate production-oriented machine-learning models.", "Build repeatable training and deployment workflows.", "Monitor model quality, drift and operational performance."]),
    ("19", "QA Automation Engineer", "QualityFirst Software", "Melbourne, VIC", [
        "Software testing and test automation are essential.",
        "You must demonstrate API testing and CI/CD integration.",
        "Python or TypeScript and agile delivery are required.",
        "Cloud platforms and container experience are preferred."],
     ["Design maintainable automated regression suites.", "Test APIs, user journeys and failure scenarios.", "Report defects with reproducible evidence and support release decisions."]),
    ("20", "Cloud Data Platform Engineer", "Southern Data Grid", "Melbourne, VIC", [
        "Cloud platforms, data warehousing and ETL pipelines are essential.",
        "You must demonstrate Python, SQL and infrastructure as code.",
        "Data governance, CI/CD and container orchestration are required.",
        "Stakeholder communication is preferred."],
     ["Operate scalable cloud data services and deployment pipelines.", "Automate infrastructure, access and environment configuration.", "Improve platform observability, cost controls and data reliability."]),
]

if __name__ == "__main__":
    HERE.mkdir(parents=True, exist_ok=True)
    for index, jd in enumerate(TECH_JDS, start=1):
        template.build(jd)
        title = jd[1]
        path = HERE / f"JD_{jd[0]}_{title.replace(' ', '_').replace('&', 'and')}.docx"
        doc = Document(path)
        footer = doc.sections[0].footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"Skill Bridge Tech JD fixture {index:02d}/10 | Synthetic, non-personal data")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("5D6B78")
        doc.save(path)
