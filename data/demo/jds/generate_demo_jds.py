"""Generate ten synthetic, upload-ready job descriptions for Skill Bridge."""

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
NAVY, TEAL, PALE, GREY = "18324B", "0B7A75", "EAF4F3", "5D6B78"

JDS = [
    ("01", "Business Operations Coordinator", "Harbour & Co Services", "Melbourne, VIC", [
        "Team operations and cross-functional coordination are essential for this role.",
        "You must prepare weekly operational reporting and maintain accurate issue logs.",
        "Strong spreadsheet analysis using Microsoft Excel or Google Sheets is required.",
        "AI tool familiarity is preferred for drafting routine updates and summaries."], ["Coordinate priorities across sales, service and fulfilment teams.", "Track actions, owners and due dates for the weekly operations meeting.", "Escalate customer and delivery issues with clear supporting evidence."]),
    ("02", "Operations Support Officer", "BrightPath Training", "Geelong, VIC", [
        "Customer service operations and schedule management are essential.",
        "The role requires operational reporting and confident stakeholder communication.",
        "Spreadsheet analysis is important for monitoring learner attendance and service volumes.",
        "Documentation management and AI tools are desirable."], ["Support day-to-day delivery of short training programs.", "Maintain schedules, attendance records and participant communications.", "Prepare weekly status reports and flag delivery risks." ]),
    ("03", "Service Delivery Coordinator", "Southern Cross Maintenance", "Dandenong, VIC", [
        "Cross-functional coordination and customer service operations are essential.",
        "You must manage escalations and maintain a current issue log.",
        "Operational reporting and stakeholder communication are required.",
        "Vendor coordination is preferred."], ["Coordinate technicians, customers and suppliers.", "Monitor open work orders and communicate service changes.", "Produce weekly service delivery and issue reports."]),
    ("04", "Workforce Operations Coordinator", "UrbanCare Support", "Melbourne, VIC", [
        "Team operations and schedule management are essential.",
        "Spreadsheet analysis and operational reporting are required.",
        "Stakeholder communication is important when resolving roster changes.",
        "Process improvement experience is preferred."], ["Maintain weekly team rosters and coverage plans.", "Analyse service demand and staffing data in Excel.", "Coordinate urgent changes with team leads and clients."]),
    ("05", "Business Support Coordinator", "Northbank Advisory", "Richmond, VIC", [
        "Project coordination and stakeholder communication are essential.",
        "You must maintain documentation and prepare management reporting.",
        "Cross-functional coordination is required across consulting teams.",
        "CRM administration and AI tool familiarity are preferred."], ["Track project actions, milestones and dependencies.", "Prepare meeting packs and concise weekly status reports.", "Maintain templates, procedures and client records."]),
    ("06", "Customer Operations Coordinator", "Loop Retail Systems", "Melbourne, VIC", [
        "Customer service operations and risk and issue management are essential.",
        "The role requires cross-functional coordination with sales and product teams.",
        "Performance analytics and spreadsheet analysis are important.",
        "CRM administration is preferred."], ["Review customer issue trends and coordinate resolutions.", "Track key performance indicators and prepare weekly dashboards.", "Maintain accurate escalations and customer follow-up records."]),
    ("07", "Project Operations Assistant", "CivicBuild Projects", "Footscray, VIC", [
        "Project coordination and documentation management are essential.",
        "You must provide operational reporting and maintain risk registers.",
        "Vendor coordination and schedule management are required.",
        "AI tool familiarity is desirable."], ["Support project plans, meeting actions and document control.", "Coordinate supplier updates and delivery schedules.", "Prepare weekly progress, risk and issue summaries."]),
    ("08", "Sales Operations Coordinator", "Greenline Equipment", "Clayton, VIC", [
        "CRM administration and operational reporting are essential.",
        "Spreadsheet analysis and stakeholder alignment are required.",
        "Cross-functional coordination with sales, finance and warehouse teams is important.",
        "Process improvement is preferred."], ["Maintain CRM data quality and weekly pipeline reporting.", "Coordinate order issues across sales, finance and dispatch.", "Document and improve recurring administrative workflows."]),
    ("09", "Program Operations Coordinator", "FutureSkills Australia", "Melbourne, VIC", [
        "Schedule management and stakeholder communication are essential.",
        "The role requires project coordination and operational reporting.",
        "Documentation management is important for program compliance.",
        "Spreadsheet analysis and AI tools are preferred."], ["Coordinate program calendars, facilitators and participant updates.", "Track milestones, attendance and delivery risks.", "Maintain program files and prepare monthly reports."]),
    ("10", "Operations & Improvement Coordinator", "Westgate Logistics", "Laverton, VIC", [
        "Process improvement and data-informed decisions are essential.",
        "You must demonstrate operational reporting and spreadsheet analysis.",
        "Team operations and cross-functional coordination are required.",
        "AI tool familiarity is preferred for workflow documentation."], ["Analyse operational trends and identify practical improvements.", "Coordinate improvement actions across warehouse and customer teams.", "Measure outcomes and report progress to operations leaders."]),
]

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)

def build(item):
    number, title, company, location, requirements, responsibilities = item
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(0.78)
    sec.header_distance = sec.footer_distance = Inches(0.4)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after in (("Heading 1",16,13,6),("Heading 2",13,10,5)):
        style = doc.styles[style_name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(TEAL)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
    bullet = doc.styles["List Bullet"]; bullet.font.name = "Calibri"; bullet.font.size = Pt(10.5); bullet.paragraph_format.left_indent = Inches(.32); bullet.paragraph_format.first_line_indent = Inches(-.18); bullet.paragraph_format.space_after = Pt(3); bullet.paragraph_format.line_spacing = 1.1

    band = doc.add_table(rows=1, cols=1); band.autofit = False; cell = band.cell(0,0); shade(cell, NAVY)
    p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(21); r.font.color.rgb = RGBColor(255,255,255)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); r = p.add_run(f"{company}  |  {location}  |  Full-time"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph("SYNTHETIC DEMO JOB DESCRIPTION - NOT A REAL VACANCY"); p.paragraph_format.space_after = Pt(7); r = p.runs[0]; r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_heading("About the role", level=1)
    doc.add_paragraph(f"{company} is seeking a practical {title} to keep information, people and operational priorities moving. This synthetic role is designed for the Skill Bridge upload and skill-mapping demo.")
    doc.add_heading("Key responsibilities", level=1)
    for text in responsibilities: doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Skills and experience", level=1)
    for text in requirements: doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("How success will be assessed", level=1)
    for text in ["Reliable follow-through on actions and deadlines.", "Clear, evidence-based communication of progress and issues.", "Accurate records and practical improvements to team workflows."]: doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Inclusive hiring note", level=1)
    doc.add_paragraph("We value transferable capability and welcome candidates whose relevant skills were gained in another country or industry. A recruiter reviews every application; this demo does not automate hiring decisions.")
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run(f"Skill Bridge JD fixture {number}/10 | Synthetic, non-personal data"); rr.font.size = Pt(8); rr.font.color.rgb = RGBColor.from_string(GREY)
    props = doc.core_properties; props.title = f"{title} - Synthetic Demo JD"; props.author = "Skill Bridge Demo Team"; props.comments = "Synthetic data only; not a real vacancy."
    path = OUT / f"JD_{number}_{title.replace(' ', '_').replace('&', 'and')}.docx"; doc.save(path); print(path)

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for jd in JDS: build(jd)
