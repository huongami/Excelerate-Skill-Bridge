"""Generate the synthetic Minh Tran CV used to test Skill Bridge uploads."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "Minh_Tran_Demo_CV.docx"

NAVY = "18324B"
TEAL = "0B7A75"
PALE = "EAF4F3"
SLATE = "44546A"
WHITE = "FFFFFF"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), TEAL)
    borders.append(bottom)
    p_pr.append(borders)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    p.add_run(text)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(2.5)

    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Cm(11.9)
    header.columns[1].width = Cm(6.0)
    remove_table_borders(header)
    left, right = header.rows[0].cells
    for cell in (left, right):
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, top=150, start=180, bottom=145, end=180)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    name = left.paragraphs[0]
    name.paragraph_format.space_after = Pt(1)
    r = name.add_run("MINH TRAN")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(23)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    title = left.add_paragraph()
    title.paragraph_format.space_after = Pt(0)
    r = title.add_run("BUSINESS OPERATIONS & CUSTOMER SERVICE")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string("8EDBD5")

    contact = right.paragraphs[0]
    contact.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    contact.paragraph_format.space_after = Pt(0)
    for index, line in enumerate([
        "Melbourne, VIC",
        "+61 400 000 000",
        "minh.tran.demo@example.com",
    ]):
        run = contact.add_run(line)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(WHITE)
        if index < 2:
            run.add_break()

    notice = doc.add_table(rows=1, cols=1)
    notice.autofit = False
    notice.columns[0].width = Cm(18.1)
    remove_table_borders(notice)
    cell = notice.cell(0, 0)
    set_cell_fill(cell, PALE)
    set_cell_margins(cell, top=70, start=130, bottom=70, end=130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("SYNTHETIC DEMO CV — NOT A REAL PERSON")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    heading(doc, "Professional profile")
    p = doc.add_paragraph(
        "International graduate and operations team lead with 3+ years of experience coordinating "
        "customer service, staffing, reporting and cross-functional issue resolution in Ho Chi Minh City. "
        "Seeking a Business Operations Coordinator role in Australia where practical team leadership and "
        "data-informed operational decision-making can transfer into an SME environment."
    )
    p.paragraph_format.space_after = Pt(2)

    heading(doc, "Core capabilities")
    skills = doc.add_table(rows=2, cols=4)
    skills.autofit = False
    remove_table_borders(skills)
    values = [
        "Team operations",
        "Customer service operations",
        "Cross-functional coordination",
        "Issue escalation",
        "Operational reporting",
        "Spreadsheet analysis",
        "Stakeholder communication",
        "Workload prioritisation",
    ]
    for idx, cell in enumerate([c for row in skills.rows for c in row.cells]):
        cell.width = Cm(4.5)
        set_cell_fill(cell, "F3F7F8")
        set_cell_margins(cell, top=55, start=85, bottom=55, end=85)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.add_run(values[idx]).bold = True

    heading(doc, "Professional experience")
    role = doc.add_table(rows=1, cols=2)
    role.autofit = False
    remove_table_borders(role)
    role.columns[0].width = Cm(13.0)
    role.columns[1].width = Cm(5.1)
    a, b = role.rows[0].cells
    set_repeat_table_header(role.rows[0])
    pa = a.paragraphs[0]
    pa.paragraph_format.space_after = Pt(0)
    rr = pa.add_run("Operations Team Lead")
    rr.bold = True
    rr.font.size = Pt(10.2)
    rr.font.color.rgb = RGBColor.from_string(NAVY)
    rr.add_break()
    rr = pa.add_run("Saigon Service Group · Ho Chi Minh City, Vietnam")
    rr.italic = True
    rr.font.size = Pt(8.8)
    pb = b.paragraphs[0]
    pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pb.paragraph_format.space_after = Pt(0)
    rr = pb.add_run("Jan 2022 – Aug 2025")
    rr.bold = True
    rr.font.size = Pt(8.8)
    bullet(doc, "Managed a 12-person operations and customer service team, coordinating daily workload and service coverage.")
    bullet(doc, "Tracked service volumes and customer issues in weekly spreadsheets to improve staffing decisions.")
    bullet(doc, "Coordinated escalations across warehouse, sales and customer support teams to keep issues moving to resolution.")
    bullet(doc, "Prepared weekly operational updates for managers, highlighting service trends, open issues and team priorities.")

    heading(doc, "Education")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Graduate Certificate in Business Analytics")
    r.bold = True
    p.add_run(" · Melbourne Institute of Business (synthetic), Melbourne · 2026")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Bachelor of Business Administration")
    r.bold = True
    p.add_run(" · Ho Chi Minh City Business College (synthetic), Vietnam · 2021")

    heading(doc, "Tools, languages & work status")
    info = doc.add_table(rows=3, cols=2)
    info.autofit = False
    remove_table_borders(info)
    info.columns[0].width = Cm(3.2)
    info.columns[1].width = Cm(14.9)
    rows = [
        ("Tools", "Microsoft Excel and Google Sheets (working knowledge); basic dashboard reporting; collaboration tools."),
        ("Languages", "Vietnamese (native); English (professional working proficiency)."),
        ("Work status", "Candidate-provided post-study work visa status for demo purposes; not verified by Skill Bridge."),
    ]
    for (label, value), row in zip(rows, info.rows):
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)
        for cell in row.cells:
            set_cell_margins(cell, top=25, start=0, bottom=25, end=80)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Skill Bridge test fixture · Synthetic, non-personal data · Generated for upload testing")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(SLATE)

    props = doc.core_properties
    props.title = "Minh Tran — Synthetic Demo CV"
    props.subject = "Skill Bridge upload test fixture"
    props.author = "Skill Bridge Demo Team"
    props.comments = "Synthetic data only; not a real person."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
