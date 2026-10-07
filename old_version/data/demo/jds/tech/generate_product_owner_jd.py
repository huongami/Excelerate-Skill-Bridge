"""Generate the synthetic Product Owner JD used in the Khoa demo video."""

import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
BASE_GENERATOR = HERE.parent / "generate_demo_jds.py"
spec = importlib.util.spec_from_file_location("skill_bridge_jd_template", BASE_GENERATOR)
template = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(template)
template.OUT = HERE

JD = (
    "21",
    "Product Owner",
    "Harbour Digital Products",
    "Melbourne, VIC",
    [
        "Product backlog management and prioritisation are essential.",
        "You must demonstrate product discovery, user research and requirements management.",
        "Stakeholder management and agile delivery are required.",
        "SQL, data analysis and dashboard development experience are preferred.",
    ],
    [
        "Own and continuously prioritise the product backlog against customer and business outcomes.",
        "Lead discovery workshops, clarify requirements and translate evidence into roadmap decisions.",
        "Partner with engineering, design and business stakeholders throughout iterative delivery.",
        "Use product and operational data to evaluate outcomes and communicate trade-offs.",
    ],
)


if __name__ == "__main__":
    template.build(JD)
    output = HERE / "JD_21_Product_Owner.docx"
    doc = Document(output)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Skill Bridge Product Owner JD fixture | Synthetic, non-personal data")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("5D6B78")
    doc.save(output)
    print(output)
